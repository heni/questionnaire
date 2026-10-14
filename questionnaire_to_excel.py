#!/usr/bin/env python3
"""Generate Excel questionnaires from YAML using only public PyPI packages.

Usage:
    python questionnaire_to_excel_openpyxl.py questionnaire.yaml output.xlsx

Dependencies:
    pip install -r requirements.txt
"""
from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

EXCEL_MAX_COL = 16384


def normalize_option(option: Any, idx: int) -> dict[str, str]:
    default_codes = ["а", "б", "в", "г", "д", "е", "ж", "з", "и", "к"]
    if isinstance(option, str):
        code = default_codes[idx] if idx < len(default_codes) else str(idx + 1)
        return {"code": code, "text": option}
    if isinstance(option, dict):
        text = str(option.get("text", "")).strip()
        if not text:
            raise ValueError(f"Option #{idx + 1} has no text")
        code = str(option.get("code") or (default_codes[idx] if idx < len(default_codes) else idx + 1))
        return {"code": code, "text": text}
    raise TypeError(f"Unsupported option type: {type(option).__name__}")


def load_config(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as f:
        cfg = yaml.safe_load(f)
    if not isinstance(cfg, dict):
        raise ValueError("YAML root must be a mapping")

    cfg.setdefault("title", "Анкета")
    cfg.setdefault("students", 40)
    cfg.setdefault("student_label", "Ученик {n}")
    cfg.setdefault("student_fields", ["ФИО", "Класс", "Возраст", "Код учащегося"])
    cfg.setdefault("statistics", True)

    students = int(cfg["students"])
    if students <= 0:
        raise ValueError("students must be positive")
    cfg["students"] = students

    fields = cfg["student_fields"]
    if not isinstance(fields, list) or not fields:
        raise ValueError("student_fields must be a non-empty list")
    cfg["student_fields"] = [str(x) for x in fields]

    sections = cfg.get("sections")
    if not isinstance(sections, list) or not sections:
        raise ValueError("sections must be a non-empty list")

    seen = set()
    normalized = []
    for sidx, section in enumerate(sections, 1):
        if not isinstance(section, dict):
            raise ValueError(f"Section #{sidx} must be a mapping")
        name = str(section.get("name") or f"Раздел {sidx}").strip()
        if len(name) > 31:
            raise ValueError(f"Excel sheet name too long: {name}")
        if name in seen:
            raise ValueError(f"Duplicate section name: {name}")
        seen.add(name)

        questions = section.get("questions")
        if not isinstance(questions, list) or not questions:
            raise ValueError(f"Section {name!r}: questions must be non-empty")

        nq = []
        for qidx, q in enumerate(questions, 1):
            if not isinstance(q, dict):
                raise ValueError(f"{name}, question #{qidx}: must be a mapping")
            text = str(q.get("text", "")).strip()
            if not text:
                raise ValueError(f"{name}, question #{qidx}: missing text")
            qid = q.get("id", qidx)
            raw = q.get("options")
            if not isinstance(raw, list) or len(raw) < 2:
                raise ValueError(f"{name}, question {qid}: at least two options required")
            opts = [normalize_option(o, i) for i, o in enumerate(raw)]
            if len(opts) > 10:
                raise ValueError(f"{name}, question {qid}: max 10 options supported")
            nq.append({"id": qid, "text": text, "options": opts})
        normalized.append({"name": name, "questions": nq})

    cfg["sections"] = normalized
    return cfg


def set_width(ws, col: int, width: float):
    ws.column_dimensions[get_column_letter(col)].width = width


def build_workbook(cfg: dict[str, Any]) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    students = cfg["students"]
    fields = cfg["student_fields"]
    title = cfg["title"]
    student_label = cfg["student_label"]
    make_stats = bool(cfg.get("statistics", True))

    first_student_col = 3  # C
    last_student_col = first_student_col + students - 1
    service_start_col = last_student_col + 2  # one blank separator
    if service_start_col + 10 > EXCEL_MAX_COL:
        raise ValueError("Too many student columns for Excel")

    fill_title = PatternFill("solid", fgColor="163A5F")
    fill_header = PatternFill("solid", fgColor="1F4E78")
    fill_sub = PatternFill("solid", fgColor="D9EAF7")
    white_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, color="FFFFFF", size=14)
    bold_font = Font(bold=True)
    thin_side = Side(style="thin", color="D3DEE5")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    section_layout = {}

    for section in cfg["sections"]:
        name = section["name"]
        questions = section["questions"]
        ws = wb.create_sheet(name)

        field_start = 2
        field_end = field_start + len(fields) - 1
        student_header_row = field_end + 1
        table_header_row = student_header_row + 1
        q_start = table_header_row + 1
        q_end = q_start + len(questions) - 1
        section_layout[name] = (q_start, q_end)

        # Title, intentionally not merged.
        ws["A1"] = f"{title} — {name}"
        ws["A1"].fill = fill_title
        ws["A1"].font = title_font
        for col in range(2, last_student_col + 1):
            ws.cell(1, col).fill = fill_title

        # Student metadata labels.
        for i, field in enumerate(fields, field_start):
            c = ws.cell(i, 1, field)
            c.fill = fill_sub
            c.font = bold_font
            c.alignment = Alignment(vertical="center")
            for col in range(first_student_col, last_student_col + 1):
                ws.cell(i, col).border = Border(bottom=Side(style="thin", color="C6D3DC"))

        # Student headers.
        for n, col in enumerate(range(first_student_col, last_student_col + 1), 1):
            c = ws.cell(student_header_row, col, student_label.format(n=n))
            c.fill = fill_sub
            c.font = bold_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Table header.
        ws.cell(table_header_row, 1, "№")
        ws.cell(table_header_row, 2, "Вопрос")
        for col in range(1, last_student_col + 1):
            c = ws.cell(table_header_row, col)
            c.fill = fill_header
            c.font = white_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Questions and service option ranges on the SAME sheet.
        max_options = max(len(q["options"]) for q in questions)
        for idx, q in enumerate(questions):
            row = q_start + idx
            ws.cell(row, 1, q["id"])
            ws.cell(row, 2, q["text"])
            ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")

            for j, opt in enumerate(q["options"]):
                ws.cell(row, service_start_col + j, opt["text"])

            start_letter = get_column_letter(service_start_col)
            end_letter = get_column_letter(service_start_col + len(q["options"]) - 1)
            formula = f"${start_letter}${row}:${end_letter}${row}"
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            dv.error = "Выберите значение из списка"
            dv.errorTitle = "Недопустимый ответ"
            dv.prompt = "Выберите один из допустимых вариантов"
            dv.promptTitle = "Ответ"
            ws.add_data_validation(dv)
            dv.add(f"C{row}:{get_column_letter(last_student_col)}{row}")

        # Borders and text wrapping.
        for row in ws.iter_rows(min_row=table_header_row, max_row=q_end, min_col=1, max_col=last_student_col):
            for c in row:
                c.border = thin_border
                c.alignment = Alignment(vertical="top", wrap_text=True)

        set_width(ws, 1, 18)
        set_width(ws, 2, 42)
        for col in range(first_student_col, last_student_col + 1):
            set_width(ws, col, 18)
        for col in range(service_start_col, service_start_col + max_options):
            set_width(ws, col, 2)
            ws.column_dimensions[get_column_letter(col)].hidden = True

        ws.freeze_panes = f"C{q_start}"
        ws.sheet_view.showGridLines = False

    if make_stats:
        for section in cfg["sections"]:
            main_name = section["name"]
            questions = section["questions"]
            stats_name = (f"Статистика {main_name}")[:31]
            if stats_name in wb.sheetnames:
                stats_name = (f"Стат {main_name}")[:31]
            ws = wb.create_sheet(stats_name)
            q_start, _ = section_layout[main_name]

            ws["A1"] = f"Статистика ответов — {main_name}"
            ws["A1"].fill = fill_title
            ws["A1"].font = title_font

            headers = ["№", "Вопрос"]
            for code in ["а", "б", "в", "г", "д", "е", "ж", "з", "и", "к"]:
                headers.extend([f"Вариант {code}", "Кол-во"])
            max_opts = max(len(q["options"]) for q in questions)
            used_headers = headers[:2 + 2 * max_opts]
            for col, h in enumerate(used_headers, 1):
                c = ws.cell(3, col, h)
                c.fill = fill_header
                c.font = white_font
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            for idx, q in enumerate(questions):
                row = 4 + idx
                main_row = q_start + idx
                ws.cell(row, 1, q["id"])
                ws.cell(row, 2, q["text"])
                for j, opt in enumerate(q["options"]):
                    text_col = 3 + 2*j
                    count_col = text_col + 1
                    ws.cell(row, text_col, opt["text"])
                    answer_range = f"'{main_name.replace(chr(39), chr(39)*2)}'!$C${main_row}:${get_column_letter(last_student_col)}${main_row}"
                    ws.cell(row, count_col, f'=COUNTIF({answer_range},{get_column_letter(text_col)}{row})')

            max_stat_col = 2 + 2 * max_opts
            for row in ws.iter_rows(min_row=3, max_row=3+len(questions), min_col=1, max_col=max_stat_col):
                for c in row:
                    c.border = thin_border
                    c.alignment = Alignment(vertical="top", wrap_text=True)

            set_width(ws, 1, 6)
            set_width(ws, 2, 42)
            for col in range(3, max_stat_col + 1):
                set_width(ws, col, 25 if col % 2 == 1 else 9)
            ws.freeze_panes = "A4"
            ws.sheet_view.showGridLines = False

    return wb


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate Excel questionnaire from YAML")
    parser.add_argument("yaml", type=Path, help="Input YAML file")
    parser.add_argument("output", type=Path, help="Output .xlsx file")
    args = parser.parse_args()

    cfg = load_config(args.yaml)
    wb = build_workbook(cfg)
    wb.save(args.output)
    print(f"Saved: {args.output}")


if __name__ == "__main__":
    main()
